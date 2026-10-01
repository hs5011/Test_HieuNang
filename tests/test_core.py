"""Kiểm thử các module lõi (không cần mạng/trình duyệt).  Chạy: .venv\\Scripts\\python -m pytest -q"""
from __future__ import annotations

import io
import xml.etree.ElementTree as ET
from pathlib import Path

import pandas as pd
import pytest

from perftool.analysis import complexity as cx
from perftool.charts.builder import build_all
from perftool.models import (AuthCapture, CapturedRequest, ModuleInfo, PageInfo, Project, RunInfo, ScenarioStep,
                             TestScenario, UseCase)
from perftool.report.docx_report import build_report
from perftool.results import metrics
from perftool.results.analyzer import analyze_run, flatten
from perftool.results.parsers import parse_jtl, parse_k6_csv
from perftool.scriptgen.generator import generate_jmeter, generate_k6
from perftool.scriptgen.profile import parse_duration, stages_for, split_vus
from perftool.uc_import import importer


# ------------------------------------------------------------------ import UC
def test_read_table_finds_header_and_maps_columns():
    csv = "Danh sách UC\n\nMã UC,Tên Use case,Phân hệ,Mô tả\nUC-1,Xem danh sách,QLCV,\"1. Mở\n2. Xem\"\nUC-2,Thêm mới,,x\n"
    df = importer.read_table(io.BytesIO(csv.encode("utf-8")), "a.csv")
    m = importer.guess_mapping(list(df.columns))
    assert m["code"] == "Mã UC" and m["name"] == "Tên Use case" and m["module"] == "Phân hệ"
    ucs = importer.to_use_cases(df, m)
    assert [u.code for u in ucs] == ["UC-1", "UC-2"]
    assert ucs[0].steps == 2
    assert ucs[1].module == "QLCV"  # ô gộp được điền xuống


def test_parse_duration_and_stages():
    assert parse_duration("1h2m3s") == 3723 and parse_duration("90") == 90 and parse_duration("5m") == 300
    p = Project(id="t")
    p.test_config.ramp_up, p.test_config.duration, p.test_config.ramp_down = "30s", "2m", "10s"
    st = stages_for(p.test_config, 20)
    assert [(s.duration_s, s.target) for s in st] == [(30, 20), (120, 20), (10, 0)]
    assert sum(split_vus(10, 3)) == 10


# ------------------------------------------------------------------ scoring
def _project() -> Project:
    p = Project(id="unit-test")
    p.login.base_url = "https://app.example.com"
    p.use_cases = [UseCase(code="UC-1", name="Thống kê giao việc", module="Công việc", steps=6),
                   UseCase(code="UC-2", name="Xem trang giới thiệu", module="Khác", steps=1)]
    reqs = [CapturedRequest(method="GET", url=f"https://app.example.com/api/stat/{i}", resource_type="xhr",
                            status=200, duration_ms=900) for i in range(8)]
    reqs.append(CapturedRequest(method="POST", url="https://app.example.com/api/tasks/save", resource_type="xhr",
                                post_data='{"a":1}'))
    p.modules = [ModuleInfo(name="Công việc", url="https://app.example.com/cv", pages=[
        PageInfo(url="https://app.example.com/cv/thong-ke", menu_text="Thống kê giao việc", module="Công việc",
                 tables=1, table_rows=60, crud_buttons=["xuất excel"], requests=reqs)])]
    return p


def test_score_and_recommend():
    p = _project()
    scores = cx.recommend(cx.score_project(p), top_n=1)
    assert scores[0].uc_code == "UC-1" and scores[0].recommended and scores[0].selected
    assert scores[0].matched_pages == ["https://app.example.com/cv/thong-ke"]
    assert not scores[1].selected
    sc = cx.build_scenario(p, scores[0])
    write = [s for s in sc.steps if s.method == "POST"]
    assert write and not write[0].enabled  # request ghi dữ liệu mặc định tắt


# ------------------------------------------------------------------ script generation
def _scenario_project(tmp_path: Path) -> Project:
    p = _project()
    p.auth = AuthCapture(login_url="https://app.example.com/api/login", login_content_type="application/json",
                         login_body_template='{"u":"{{USERNAME}}","p":"{{PASSWORD}}"}', token_json_path="data.token")
    p.scenarios = [TestScenario(uc_code="UC-1", uc_name="Thống kê <&> giao việc", steps=[
        ScenarioStep(name="Mở trang", url="https://app.example.com/cv?x=1&y=2"),
        ScenarioStep(name="Tìm", method="POST", url="https://app.example.com/api/search", body='{"k":"<a>"}',
                     content_type="application/json", think_time_s=1)])]
    return p


def test_generate_k6_and_jmeter(tmp_path):
    p = _scenario_project(tmp_path)
    js = generate_k6(p, p.scenarios, tmp_path / "s.js").read_text(encoding="utf-8")
    assert "noCookiesReset: true" in js and "{{USERNAME}}" in js and "p(95)<3000" in js
    jmx = generate_jmeter(p, p.scenarios, tmp_path / "p.jmx")
    root = ET.parse(jmx).getroot()  # XML hợp lệ (đã escape ký tự đặc biệt)
    names = [e.get("testname") for e in root.iter("HTTPSamplerProxy")]
    assert "UC-1 :: LOGIN" in names and "UC-1 :: Tìm" in names
    assert any(e.text == "$.data.token" for e in root.iter("stringProp"))


# ------------------------------------------------------------------ results
K6_CSV = """metric_name,timestamp,metric_value,check,error,error_code,expected_response,group,method,name,proto,scenario,service,status,subproto,tls_version,url,extra_tags,metadata
vus,1000,2,,,,,,,,,,,,,,,,
http_req_duration,1000,100,,,,true,,GET,UC-1 :: A,HTTP/1.1,UC_1,,200,,,u,uc=UC-1,
http_req_duration,1001,300,,,,true,,GET,UC-1 :: B,HTTP/1.1,UC_1,,200,,,u,uc=UC-1,
http_req_duration,1002,500,,,,false,,GET,UC-1 :: B,HTTP/1.1,UC_1,,500,,,u,uc=UC-1,
http_req_duration,1002,50,,,,true,,POST,UC-1 :: LOGIN,HTTP/1.1,UC_1,,200,,,u,uc=UC-1,
vus,1002,4,,,,,,,,,,,,,,,,
"""
JTL = """timeStamp,elapsed,label,responseCode,responseMessage,threadName,dataType,success,failureMessage,bytes,sentBytes,grpThreads,allThreads,URL,Latency,IdleTime,Connect
1000000,100,UC-1 :: A,200,OK,t1,text,true,,10,1,1,1,u,90,0,1
1001000,300,UC-1 :: B,200,OK,t1,text,true,,10,1,2,2,u,90,0,1
1002000,400,UC-1 :: TRANSACTION,200,Number of samples,t1,,true,,20,2,2,2,,0,0,0
1002000,900,UC-1 :: B,503,ERR,t2,text,false,x,10,1,2,2,u,90,0,1
"""


def test_parsers_and_metrics(tmp_path):
    (tmp_path / "raw.csv").write_text(K6_CSV)
    df, vus = parse_k6_csv(tmp_path / "raw.csv")
    assert len(df) == 4 and df["is_login"].sum() == 1 and vus["vus"].max() == 4
    s = metrics.summarize(df[~df["is_login"]], vus)
    assert s["samples"] == 3 and s["errors"] == 1 and s["max"] == 500 and s["max_vus"] == 4
    ev = metrics.evaluate(s, 400, 0.01)
    assert not ev["passed"] and len(ev["reasons"]) == 2

    (tmp_path / "r.jtl").write_text(JTL)
    df2, vus2 = parse_jtl(tmp_path / "r.jtl")
    assert len(df2) == 3  # bỏ dòng TRANSACTION
    assert metrics.summarize(df2)["error_rate"] == pytest.approx(1 / 3)


def test_analyze_charts_report(tmp_path, monkeypatch):
    import perftool.storage as storage
    monkeypatch.setattr(storage, "PROJECTS_DIR", tmp_path / "projects")
    import perftool.charts.builder  # noqa: F401
    p = _scenario_project(tmp_path)
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    (run_dir / "raw.csv").write_text(K6_CSV)
    (run_dir / "script.js").write_text("//")
    run = RunInfo(run_id="r1", tool="k6", uc_code="UC-1", status="done", script_file=str(run_dir / "script.js"),
                  raw_file=str(run_dir / "raw.csv"), config=p.test_config.model_dump())
    p.runs = [run]
    res = analyze_run(p, run, force=True)
    rows = flatten([res])
    assert rows and rows[0]["uc_code"] == "UC-1" and rows[0]["insights"]
    per_row, cmp = build_all(p, rows, tmp_path / "charts")
    assert Path(per_row["r1|UC-1"]["percentile"]).exists()
    out = build_report(p, rows, per_row, cmp, tmp_path / "bc.docx")
    from docx import Document
    text = "\n".join(par.text for par in Document(str(out)).paragraphs)
    assert "BÁO CÁO KẾT QUẢ KIỂM THỬ HIỆU NĂNG" in text and "TỔNG HỢP VÀ KẾT LUẬN" in text
    assert isinstance(pd.DataFrame(rows), pd.DataFrame)


def test_overview_section(tmp_path, monkeypatch):
    """Nhiều kịch bản -> báo cáo có mục tổng quan + biểu đồ tổng quát."""
    import perftool.storage as storage
    from docx import Document
    from perftool.charts.builder import build_overview
    from perftool.results.insights import overview_insights
    monkeypatch.setattr(storage, "PROJECTS_DIR", tmp_path / "projects")
    p = _scenario_project(tmp_path)
    rows = []
    for tool, typ in (("k6", "smoke"), ("jmeter", "load")):
        d = tmp_path / f"run_{tool}"
        d.mkdir()
        if tool == "k6":
            (d / "raw.csv").write_text(K6_CSV)
        else:
            (d / "raw.csv").write_text(JTL)
        (d / "script.js").write_text("//")
        cfg = p.test_config.model_copy(update={"scenario_type": typ}).model_dump()
        run = RunInfo(run_id=f"r_{tool}", tool=tool, uc_code="UC-1", scenario_type=typ, status="done",
                      script_file=str(d / "script.js"), raw_file=str(d / "raw.csv"), config=cfg)
        rows += flatten([analyze_run(p, run, force=True)])
    assert len(rows) == 2
    notes = overview_insights(rows)
    assert notes and "2 lượt kiểm thử" in notes[0] and "Smoke Test, Load Test" in notes[0]
    ov = build_overview(rows, tmp_path / "ov")
    assert {"vus", "p95", "throughput", "comparison"} <= set(ov)
    per_row, cmp = build_all(p, rows, tmp_path / "charts")
    out = build_report(p, rows, per_row, cmp, tmp_path / "bc.docx", {"overview": True, "overview_charts": ov})
    heads = [x.text for x in Document(str(out)).paragraphs if x.style.name.startswith("Heading")]
    assert any("TỔNG QUAN CÁC KỊCH BẢN" in h for h in heads)
    assert any(h.startswith("2.2.") for h in heads)       # chi tiết được đánh số sau mục tổng quan


def test_report_grouped_by_module(tmp_path, monkeypatch):
    """2 phân hệ × 2 UC -> báo cáo có mục tổng hợp theo phân hệ, bảng nhóm theo phân hệ, kết luận theo phân hệ."""
    import perftool.storage as storage
    from docx import Document
    from perftool.charts.builder import build_overview
    from perftool.results.insights import module_summaries
    monkeypatch.setattr(storage, "PROJECTS_DIR", tmp_path / "projects")
    p = _scenario_project(tmp_path)
    base = flatten([analyze_run(p, RunInfo(run_id="r0", tool="k6", uc_code="UC-1", scenario_type="load", status="done",
                                           script_file=str(_run_dir(tmp_path, "r0") / "script.js"),
                                           raw_file=str(_run_dir(tmp_path, "r0") / "raw.csv"),
                                           config=p.test_config.model_dump()), force=True)])[0]
    rows = []
    for mod, codes in (("Phân hệ A", ("UC-A1", "UC-A2")), ("Phân hệ B", ("UC-B1", "UC-B2"))):
        for c in codes:
            passed = c != "UC-B2"
            rows.append({**base, "uc_code": c, "uc_name": f"Tên {c}", "module": mod, "run_id": f"run_{c}",
                         "evaluation": {**base["evaluation"], "passed": passed, "p95_ok": True, "err_ok": passed,
                                        "reasons": [] if passed else ["tỷ lệ lỗi 5.00% > ngưỡng 1.00%"]}})
    sums = {m["module"]: m for m in module_summaries(rows)}
    assert sums["Phân hệ A"]["verdict"] == "Đạt" and sums["Phân hệ A"]["ucs"] == ["UC-A1", "UC-A2"]
    assert sums["Phân hệ B"]["verdict"] == "Đạt một phần" and sums["Phân hệ B"]["failed_ucs"] == ["UC-B2"]
    ov = build_overview(rows, tmp_path / "ov")
    assert "modules" in ov
    out = build_report(p, rows, {}, "", tmp_path / "bc.docx", {"overview": False, "overview_charts": ov, "modules": True})
    d = Document(str(out))
    heads = [x.text for x in d.paragraphs if x.style.name.startswith("Heading")]
    assert any("TỔNG HỢP KẾT QUẢ THEO PHÂN HỆ" in h for h in heads)
    text = "\n".join(x.text for x in d.paragraphs)
    assert "Kết luận theo phân hệ" in text and "Phân hệ B: Đạt một phần (1/2" in text
    cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
    assert any(c.startswith("1. Phân hệ A – 2/2 kịch bản đạt") for c in cells)      # dòng nhóm trong bảng tổng hợp


def _run_dir(tmp_path, name):
    d = tmp_path / name
    if not d.exists():
        d.mkdir()
        (d / "raw.csv").write_text(K6_CSV)
        (d / "script.js").write_text("//")
    return d

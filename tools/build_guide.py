"""Chuyển HUONG_DAN_SU_DUNG.md -> HUONG_DAN_SU_DUNG.docx (dùng style của file mẫu báo cáo).

    .venv\\Scripts\\python tools\\build_guide.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor  # noqa: E402

from perftool.report.docx_report import FONT_PT, ReportBuilder  # noqa: E402

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|<https?://[^>]+>|\*[^*]+\*)")


def add_inline(p, text: str, size: float = FONT_PT, bold: bool = False) -> None:
    """Hỗ trợ **đậm**, *nghiêng*, `code`, <link>."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.color.rgb = RGBColor(0x9C, 0x1C, 0x1C)
        elif part.startswith("<http"):
            r = p.add_run(part[1:-1])
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x9C)
            r.underline = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1])
            r.italic = True
        else:
            r = p.add_run(part)
            r.bold = bold
        r.font.size = Pt(size)


def shade(p, fill: str) -> None:
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def table_widths(n: int) -> list[int]:
    total = 9600
    if n == 2:
        return [3000, 6600]
    if n == 3:
        return [2200, 2600, 4800]
    return [total // n] * n


def convert(md: Path, out: Path) -> Path:
    rb = ReportBuilder()
    lines = md.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("```"):                                  # khối code
            i += 1
            block = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            for b in block:
                p = rb.doc.add_paragraph()
                r = p.add_run(b if b else " ")
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                shade(p, "F3F4F6")
            rb.para(after=60)
            continue
        if s.startswith("|"):                                    # bảng
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            header, body = rows[0], rows[1:]
            t = rb.table(header, [[""] * len(header) for _ in body], table_widths(len(header)))
            for ri, row in enumerate(body, start=1):
                for ci, val in enumerate(row[:len(header)]):
                    cell = t.rows[ri].cells[ci]
                    cell.paragraphs[0].text = ""
                    add_inline(cell.paragraphs[0], val, size=11)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:                                                    # tiêu đề
            level = len(m.group(1))
            if level == 1:
                rb.para(m.group(2), bold=True, size=17, align=WD_ALIGN_PARAGRAPH.CENTER, after=200)
            else:
                rb.heading(m.group(2).replace("`", ""), level - 1)
            i += 1
            continue
        if s.startswith(">"):                                    # ghi chú
            p = rb.doc.add_paragraph()
            add_inline(p, "Lưu ý: " + s.lstrip("> ").strip(), size=11)
            shade(p, "FFF4E5")
            p.paragraph_format.space_after = Pt(6)
            i += 1
            continue
        mb = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)", line)
        if mb:                                                   # danh sách
            indent = len(mb.group(1)) // 2
            marker = "•" if mb.group(2) in ("-", "*") else mb.group(2)
            p = rb.doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18 + 18 * indent)
            p.paragraph_format.first_line_indent = Pt(-12)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, f"{marker} {mb.group(3)}")
            i += 1
            continue
        p = rb.doc.add_paragraph()                               # đoạn văn
        add_inline(p, s)
        p.paragraph_format.space_after = Pt(6)
        i += 1
    return rb.save(out)


if __name__ == "__main__":
    src = ROOT / "HUONG_DAN_SU_DUNG.md"
    print(convert(src, src.with_suffix(".docx")))
